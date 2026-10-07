"""Numeric Excel, PowerPoint and plain-English documents from computed outputs."""

import argparse
import json

import pandas as pd
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt

from .data import output_root
from .index_engine import metrics


def _csv(name):
    return pd.read_csv(output_root() / "reports" / name)


def summary(results, bench, cfg, secondary=None):
    rows = []
    for name, result in results.items():
        for mode, column in [("net", "level"), ("gross", "gross_level")]:
            levels = pd.DataFrame(
                {
                    "level": result.levels[column],
                    "return_": result.levels["return_" if mode == "net" else "gross_return"],
                }
            )
            values = metrics(levels, bench.return_, cfg)
            rows.append(
                dict(
                    index=name,
                    mode=mode,
                    performance_benchmark=bench.attrs.get("label", cfg["benchmark"]),
                    start=levels.index.min(),
                    end=levels.index.max(),
                    rebalances=len(result.turnover),
                    average_turnover=result.turnover.turnover.mean(),
                    final_level=levels.level.iloc[-1],
                    **values,
                )
            )
    rows.append(
        dict(
            index=bench.attrs.get("label", cfg["benchmark"]),
            mode="gross",
            start=bench.index.min(),
            end=bench.index.max(),
            final_level=bench.level.iloc[-1],
            **metrics(bench, bench.return_, cfg),
        )
    )
    if secondary is not None and secondary.attrs.get("label") != bench.attrs.get("label"):
        rows.append(
            dict(
                index=secondary.attrs["label"],
                mode="gross",
                start=secondary.index.min(),
                end=secondary.index.max(),
                final_level=secondary.level.iloc[-1],
                **metrics(secondary, secondary.return_, cfg),
            )
        )
    return pd.DataFrame(rows)


def format_workbook(book):
    percentages = {
        "strike_rate",
        "zero_rate",
        "rate",
        "vol2",
        "q",
        "vol_change",
        "spot_shock",
        "first_call_probability",
        "capital_loss_prob",
        "CAGR",
        "vol",
        "max_drawdown",
        "tracking_error",
        "average_turnover",
        "return_",
        "gross_return",
        "cost",
        "weight",
        "before",
        "after",
        "trade",
        "turnover",
        "estimated_cost",
        "allocation",
        "selection",
        "interaction",
        "portfolio_gross",
        "portfolio_net",
        "benchmark",
        "active_gross",
        "active_net",
        "cost_effect",
        "pnl",
        "worst_day",
        "w_p",
        "w_b",
        "R_p",
        "R_b",
        "top3",
        "HHI",
        "contribution",
        "arithmetic_active",
        "compounded_active",
        "compounding_residual",
        "total_cost_effect",
    }
    for sheet in book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        headers = {cell.column: str(cell.value) for cell in sheet[1]}
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = (
                        "0.00%"
                        if (
                            headers[cell.column] in percentages
                            or (
                                sheet.title == "Weights"
                                and headers[cell.column] not in {"date", "index"}
                            )
                        )
                        else "0"
                        if headers[cell.column] in {"rebalances", "observations", "count"}
                        else "0.0000"
                    )
                elif isinstance(cell.value, (pd.Timestamp,)) or hasattr(cell.value, "year"):
                    cell.number_format = "yyyy-mm-dd"
        for cells in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in list(cells)[:100])
            sheet.column_dimensions[cells[0].column_letter].width = min(max(width + 2, 14), 36)


def charts(results, bench, secondary=None):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    for name, result in results.items():
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.plot(result.levels.index, result.levels.level, label=f"{name} net")
        ax.plot(bench.index, bench.level, label=bench.attrs.get("label", "Benchmark"))
        if secondary is not None:
            ax.plot(secondary.index, secondary.level, label="Equal-weight sample proxy")
        ax.set(ylabel="Level", title=f"{name}: sample universe, unverified membership")
        ax.legend()
        ax.grid(alpha=0.2)
        fig.tight_layout()
        fig.savefig(output_root() / "reports" / f"{name}_level.png", dpi=140)
        plt.close(fig)


def _text(slide, text, x, y, w, h, size=15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(size)
    return box


def cfg_kind(name):
    return "INR multi-asset" if name == "MULTI" else "Equity"


def factsheet(results, bench, stats):
    deck = Presentation()
    deck.slide_width = Inches(12)
    deck.slide_height = Inches(7.5)
    for name, result in results.items():
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _text(slide, f"{name} | {cfg_kind(name)} index sample", 0.4, 0.2, 11, 0.5, 24)
        slide.shapes.add_picture(
            str(output_root() / "reports" / f"{name}_level.png"),
            Inches(0.4),
            Inches(1),
            width=Inches(7),
        )
        row = stats[(stats["index"] == name) & (stats["mode"] == "net")].iloc[0]
        table = slide.shapes.add_table(6, 2, Inches(7.6), Inches(1), Inches(4), Inches(2.5)).table
        for i, field in enumerate(
            ["CAGR", "vol", "Sharpe", "max_drawdown", "tracking_error", "average_turnover"]
        ):
            table.cell(i, 0).text = field
            table.cell(i, 1).text = (
                f"{row[field]:.2f}" if field == "Sharpe" else f"{row[field]:.2%}"
            )
            for cell in table.rows[i].cells:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(13)
        stock = _csv(f"{name}_composition.csv").sort_values("weight", ascending=False).head(5)
        _text(
            slide,
            "Latest closing constituents\n"
            + "\n".join(f"{r.symbol}: {r.weight:.1%}" for r in stock.itertuples()),
            7.6,
            3.8,
            4,
            1.7,
            13,
        )
        contributor = _csv(f"{name}_contributors.csv").iloc[0]
        commentary = (
            f"Net CAGR {row.CAGR:.2%}; comparison vs "
            f"{bench.attrs.get('label', 'Benchmark')}.\nLargest linked gross "
            f"contributor: {contributor.symbol}, "
            f"{contributor.contribution:.2%}.\nHistorical max drawdown "
            f"{row.max_drawdown:.2%}; future losses can exceed sample."
        )
        _text(slide, commentary, 0.4, 5.4, 11, 1.2, 15)
        _text(
            slide,
            f"{row.start.date()} to {row.end.date()} | {int(row.rebalances)} "
            f"rebalances | Source: mdq; adjustments unverified",
            0.4,
            6.8,
            11,
            0.3,
            10,
        )
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide, "Current-position stress replay", 0.4, 0.3, 11, 0.5, 24)
    lines = []
    for name in results:
        for r in _csv(f"{name}_historical_stress.csv").itertuples():
            lines.append(
                f"{name} {r.start[:10]}—{r.end[:10]}: P&L {r.pnl:.2%}, drawdown "
                f"{r.max_drawdown:.2%}, worst day {r.worst_day:.2%}"
            )
    _text(slide, "\n".join(lines), 0.5, 1.2, 11, 4.5, 17)
    _text(
        slide,
        "Actual daily historical weights replayed. Worst windows selected "
        "from sample proxy; historical betas are estimates.",
        0.5,
        6,
        11,
        1,
        15,
    )
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide, "Attribution: equity vs proxy; MULTI vs fixed sleeves", 0.4, 0.3, 11, 0.5, 24)
    lines = []
    for name in results:
        a = _csv(f"{name}_attribution.csv")
        res = _csv(f"{name}_compounding_bridge.csv").iloc[0]
        sums = a[["allocation", "selection", "interaction"]].sum()
        lines.append(
            f"{name}: allocation {sums.allocation:.2%}, selection "
            f"{sums.selection:.2%}, interaction {sums.interaction:.2%}"
        )
        lines.append(
            f"Compounded active {res.compounded_active:.2%}; "
            f"arithmetic-to-compound residual {res.compounding_residual:.2%}; "
            f"net cost effect {res.total_cost_effect:.2%}."
        )
    _text(slide, "\n\n".join(lines), 0.5, 1.5, 11, 4, 18)
    _text(
        slide,
        "Static illustrative sectors. Monthly effects reconcile; sums "
        "require explicit compounding bridge.",
        0.5,
        6,
        11,
        1,
        15,
    )
    deck.save(output_root() / "reports/factsheet.pptx")


def documents(results, bench, stats, cfg, log):
    from .monitor import exception_explanations

    explained = exception_explanations(log, results, cfg)
    counts = log.groupby(["check_name", "severity"]).size().rename("count").reset_index()
    from .drafts import table

    review = [
        "# Computed monitor review",
        "Source: monitor_log.csv and index levels/contributions.",
        table(counts),
        table(explained),
    ]
    (output_root() / "docs/monitor_review.md").write_text("\n\n".join(review) + "\n")
    methods = [
        "# Index descriptions",
        f"Universe: {cfg['universe_label']}. Actual available names: {len(cfg['universe'])}.",
        f"Performance/beta benchmark: {bench.attrs.get('label', cfg['benchmark'])}.",
        "Brinson benchmark: monthly equal-weight sample proxy; no Nifty sector weights.",
    ]
    launch = [
        "# Repeat Index Launch Form — educational draft",
        "Approval fields deliberately blank; no launch approval implied.",
    ]
    comments = [
        "# Performance commentary",
        "This is a sample-universe backtest, not a full Nifty 50 strategy "
        "or an investment recommendation.",
    ]
    changes = [
        "# Example methodology change notes",
        "Simulation of modifications; no live product or implementation approval implied.",
    ]
    for name, result in results.items():
        rule = cfg["indices"][name]
        r = stats[(stats["index"] == name) & (stats["mode"] == "net")].iloc[0]
        description = (
            f"Rank {cfg['momentum_months']}-{cfg['skip_months']} calendar-month momentum descending"
            if rule["rule"] == "momentum"
            else f"Rank trailing {cfg['vol_days']}-session simple-return volatility ascending"
        )
        if name == "MULTI":
            description = f"Fixed asset sleeves {cfg['multi']['targets']}; monthly targets"
        methods += [
            f"## {name}",
            f"Objective: transparent {rule['rule']} exposure. Selection: "
            f"{description}; {rule['n']} holdings. Weight: "
            + (f"fixed targets {cfg['multi']['targets']}" if name == "MULTI" else "equal")
            + " at each rebalance, drift between rebalances.",
            f"Decision: last observed session each month; effective: following "
            f"observed session. Base date {r.start.date()}, base level "
            f"{cfg['base_level']}. Cost {cfg['cost_bps']} bps per side on "
            f"absolute traded weight, including initial purchases. Calculate "
            f"return from beginning weights; cost reduces capital before return.",
            f"Sample: {r.start.date()} to {r.end.date()}, {int(r.rebalances)} "
            f"rebalances. Source: existing mdq SQLite "
            f"{cfg['price_table']}.{cfg['price_field']}. Missing, invalid and "
            f"overnight returns beyond the configured calculation gate block calculation.",
        ]
        launch += [
            f"## {name}",
            f"""Index name: {name}
Methodology reference: index_methodology.md
Universe: {cfg["universe_label"]}
Selection: {description}; {rule["n"]} names
Base date: {r.start.date()}
Base level: {cfg["base_level"]}
Schedule: monthly, next-session execution
Cost: {cfg["cost_bps"]} bps per side
Performance benchmark: {bench.attrs.get("label", cfg["benchmark"])}
Attribution benchmark: monthly equal-weight sample proxy
Data: {cfg["database"]}; {cfg["price_field"]}
Requested by: ______
Reviewed by: ______
Approved by: ______
Approval date: ______
Publication date: ______""",
        ]
        a = _csv(f"{name}_attribution.csv")
        sector = (
            a.groupby("sector")[["allocation", "selection", "interaction"]]
            .sum()
            .sum(axis=1)
            .sort_values(ascending=False)
        )
        contributor = _csv(f"{name}_contributors.csv").iloc[0]
        bridge = _csv(f"{name}_compounding_bridge.csv").iloc[0]
        comments += [
            f"## {name}",
            f"Between {r.start.date()} and {r.end.date()}, annualized net growth "
            f"was {r.CAGR:.2%}, after modeled trading costs. The monthly "
            f"equal-weight sample proxy grew at "
            f"{stats.loc[stats['index'] == 'Equal-weight sample proxy', 'CAGR'].iloc[0]:.2%} "
            f"annually. These results describe {int(r.rebalances)} rebalances, "
            f"not a forecast.",
            f"The largest linked gross stock contribution was "
            f"{contributor.symbol}, {contributor.contribution:.2%} of starting "
            f"capital. {sector.index[0]} had the largest summed monthly sector "
            f"effect ({sector.iloc[0]:.2%}); this includes allocation, selection "
            f"and interaction. The compounding bridge is "
            f"{bridge.compounding_residual:.2%}; monthly effect sums should not "
            f"be read as compounded excess performance.",
            f"One risk: concentrated equity exposure experienced a "
            f"{r.max_drawdown:.2%} maximum historical drawdown. One limitation: "
            f"membership and corporate-action adjustments are unverified; "
            f"current holdings and static sector groups can flatter "
            f"retrospective results.",
        ]
        demo = _csv(f"{name}_modifications.csv") if name != "MULTI" else pd.DataFrame()
        changes += [
            f"## {name}",
            "Before/after independently rebuilt; effective historical start is "
            "unchanged. Review and approval: ______.",
            demo.to_string(index=False),
        ]
    methods.append(
        "Disclaimer: educational sample, static universe/sectors, "
        "unverified adjustments, flat modeled costs; no live calculation, "
        "dealer calibration or vendor feed. Historical returns do not "
        "establish future returns."
    )
    for filename, lines in [
        ("index_methodology.md", methods),
        ("index_launch_form.md", launch),
        ("commentary.md", comments),
        ("index_change_note_example.md", changes),
    ]:
        (output_root() / "docs" / filename).write_text("\n\n".join(lines) + "\n")
    (output_root() / "docs/index_change_note_template.md").write_text(
        """# Index change note template

Index: ______
Proposed effective date: ______
Reason: ______
Old/new rule and config diff: ______
Backtest comparison: ______
Turnover and cost impact: ______
Data/calendar review: ______
Reviewer: ______
Approval: ______
Client notification: ______
"""
    )
    status = {
        "backtest_start": str(bench.index.min().date()),
        "backtest_end": str(bench.index.max().date()),
        "source_start": cfg.get("source_start"),
        "source_end": cfg.get("source_end"),
        "rebalances": {n: len(r.turnover) for n, r in results.items()},
        "universe_verified": "unverified",
        "full_nifty50": "unverified",
        "author_primer_complete": "unverified",
        "notes_published": "unverified",
        "not_done": cfg.get("not_done", []),
        "performance_benchmark": bench.attrs.get("label", cfg["benchmark"]),
        "attribution_benchmark": "Equal-weight sample proxy",
        "monitor_exception_rows": len(log),
        "fast": cfg.get("fast", False),
        "mc_paths": cfg["mc_paths"],
        "seed": cfg["seed"],
    }
    (output_root() / "reports/run_manifest.json").write_text(json.dumps(status, indent=2) + "\n")


def generate(results, bench, cfg, log, secondary=None):
    stats = summary(results, bench, cfg, secondary)
    stats.to_csv(output_root() / "reports/summary.csv", index=False)

    def combine(suffix):
        return pd.concat([_csv(f"{name}_{suffix}.csv") for name in results], ignore_index=True)

    monthly = combine("monthly")
    monthly["month"] = pd.to_datetime(monthly.month)
    sheets = {
        "Summary": stats,
        "Bootstrap": _csv("bootstrap.csv"),
        "Parameter sensitivity": _csv("parameter_sensitivity.csv"),
        "Drawdown": _csv("drawdown.csv"),
        "Drawdown detractors": _csv("drawdown_top5.csv"),
        "Monthly returns": monthly,
        "Weights": pd.concat(
            [r.weights.reset_index().assign(index=n) for n, r in results.items()], ignore_index=True
        ),
        "Turnover": pd.concat(
            [r.turnover.assign(index=n) for n, r in results.items()], ignore_index=True
        ),
        "Attribution": combine("attribution"),
        "Attribution bridge": combine("compounding_bridge"),
        "Contributors": combine("contributors"),
        "Stress": pd.concat(
            [combine("historical_stress"), combine("hypothetical_stress")], ignore_index=True
        ),
        "Composition": combine("composition"),
        "Sector weights": combine("sector_weights"),
        "Concentration": combine("concentration"),
        "Monitor exceptions": log,
        "Rebalance proposals": combine("rebalance_report"),
    }
    with pd.ExcelWriter(output_root() / "reports/performance.xlsx", engine="openpyxl") as writer:
        for name, table in sheets.items():
            table = table.copy()
            for col in ["date", "month", "start", "end", "asof", "decision_date", "effective_date"]:
                if col in table:
                    table[col] = pd.to_datetime(table[col])
            table.to_excel(writer, sheet_name=name, index=False)
        grid = pd.read_csv(output_root() / "reports/sensitivity_grid.csv", index_col=0)
        grid.columns = grid.columns.astype(float)
        grid.to_excel(writer, sheet_name="Sensitivity grid")
        format_workbook(writer.book)
        for cell in list(writer.book["Sensitivity grid"][1])[1:]:
            cell.number_format = "0.0%"
        for cells in writer.book["Sensitivity grid"].iter_rows(min_row=2, min_col=1, max_col=1):
            cells[0].number_format = "0.0%"
    charts(results, bench, secondary)
    factsheet(results, bench, stats)
    documents(results, bench, stats, cfg, log)
    return stats


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--period", choices=["monthly"], required=True)
    parser.add_argument("--asof", required=True)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--fast", action="store_true")
    mode.add_argument("--full", action="store_true")
    parser.add_argument("--config")
    parser.add_argument("--output-dir")
    args = parser.parse_args()
    asof = pd.Timestamp(args.asof)
    if pd.isna(asof):
        raise ValueError("Valid asof date required")
    from indexkit.data import config, load

    cfg = config(args.config)
    frame, _ = load(cfg)
    if asof > frame.date.max():
        raise ValueError("As-of after last supplied observation")
    cfg["end"] = str(asof.date())
    from run_all import run

    run(not args.full, cfg, args.output_dir)
    print(f"Monthly report rebuilt through last supplied session on/before {asof.date()}")


if __name__ == "__main__":
    main()
